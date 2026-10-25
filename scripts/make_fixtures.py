#!/usr/bin/env python3
"""Crear DOCX y ODT externos usando python-docx y odfpy."""
from pathlib import Path
from docx import Document as WordDocument
from odf.opendocument import OpenDocumentText
from odf.text import P, Span
from odf.style import Style, TextProperties
from odf.dc import Title

root = Path(__file__).resolve().parents[1] / "fixtures"
root.mkdir(parents=True, exist_ok=True)

word = WordDocument()
word.core_properties.title = "Documento externo de prueba"
word.add_paragraph("Cabecera sin modificar")
p = word.add_paragraph()
p.add_run("Texto").bold = True
p.add_run(" de ejemplo para comprobar el formato.")
word.save(root / "externo.docx")

odt = OpenDocumentText()
odt.meta.addElement(Title(text="Documento externo de prueba"))
strong = Style(name="StrongProbe", family="text")
strong.addElement(TextProperties(fontweight="bold"))
odt.styles.addElement(strong)
odt.text.addElement(P(text="Cabecera sin modificar"))
p = P()
p.addElement(Span(stylename=strong, text="Texto"))
p.addText(" de ejemplo para comprobar el formato.")
odt.text.addElement(p)
odt.save(str(root / "externo.odt"))
